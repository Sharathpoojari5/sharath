"""Builds the 'Small Business Toolkit India' Excel templates (works in Excel + Google Sheets)."""
import os, zipfile
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import BarChart, Reference

OUT = os.path.join(os.path.dirname(__file__), "dist")
os.makedirs(OUT, exist_ok=True)
HEAD = PatternFill("solid", fgColor="1F4E78")
INPUT = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="BBBBBB")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

def hdr(ws, row, cols):
    for i, t in enumerate(cols, 1):
        c = ws.cell(row=row, column=i, value=t)
        c.font = Font(bold=True, color="FFFFFF"); c.fill = HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = BOX

def widths(ws, ws_w):
    for i, w in enumerate(ws_w, 1):
        ws.column_dimensions[chr(64 + i)].width = w

def instructions(wb, title, lines):
    ws = wb.create_sheet("READ ME", 0)
    ws["A1"] = title; ws["A1"].font = Font(bold=True, size=16, color="1F4E78")
    for i, l in enumerate(lines, 3):
        ws.cell(row=i, column=1, value=l).alignment = Alignment(wrap_text=True)
    ws.cell(row=len(lines) + 4, column=1, value="Yellow cells = you type. Everything else calculates automatically.").font = Font(italic=True)
    ws.column_dimensions["A"].width = 100

# 1 ---------------- GST invoice
def gst_invoice():
    wb = Workbook(); ws = wb.active; ws.title = "Invoice"
    ws["A1"] = "TAX INVOICE"; ws["A1"].font = Font(bold=True, size=18, color="1F4E78")
    meta = [("Your business name", ""), ("Your address", ""), ("Your GSTIN", ""), ("Invoice no.", "INV-001"),
            ("Invoice date", ""), ("Bill to (customer)", ""), ("Customer GSTIN", ""), ("Customer state code", ""),
            ("Your state code", "")]
    for i, (k, v) in enumerate(meta, 3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        c = ws.cell(row=i, column=2, value=v); c.fill = INPUT; c.border = BOX
    r0 = 13
    hdr(ws, r0, ["#", "Item / service", "HSN/SAC", "Qty", "Rate (₹)", "GST %", "Taxable (₹)", "CGST (₹)", "SGST (₹)", "IGST (₹)", "Line total (₹)"])
    gst_dv = DataValidation(type="list", formula1='"0,5,12,18,28"', allow_blank=True); ws.add_data_validation(gst_dv)
    for n in range(1, 16):
        r = r0 + n
        ws.cell(row=r, column=1, value=n)
        for c in (2, 3, 4, 5, 6):
            ws.cell(row=r, column=c).fill = INPUT
        gst_dv.add(ws.cell(row=r, column=6))
        ws.cell(row=r, column=7, value=f'=IF(D{r}="","",D{r}*E{r})')
        intra = "$B$10=$B$11"
        ws.cell(row=r, column=8, value=f'=IF(G{r}="","",IF({intra},G{r}*F{r}/100/2,0))')
        ws.cell(row=r, column=9, value=f'=IF(G{r}="","",IF({intra},G{r}*F{r}/100/2,0))')
        ws.cell(row=r, column=10, value=f'=IF(G{r}="","",IF({intra},0,G{r}*F{r}/100))')
        ws.cell(row=r, column=11, value=f'=IF(G{r}="","",G{r}+H{r}+I{r}+J{r})')
        for c in range(1, 12): ws.cell(row=r, column=c).border = BOX
    t = r0 + 16
    ws.cell(row=t, column=6, value="TOTAL").font = Font(bold=True)
    for c, col in zip(range(7, 12), "GHIJK"):
        x = ws.cell(row=t, column=c, value=f"=SUM({col}{r0+1}:{col}{r0+15})"); x.font = Font(bold=True); x.border = BOX
    ws.cell(row=t + 2, column=1, value="Place of supply rule: same state code = CGST+SGST, different = IGST (auto).").font = Font(italic=True)
    ws.cell(row=t + 3, column=1, value="Bank details / UPI ID:").font = Font(bold=True)
    ws.cell(row=t + 3, column=2).fill = INPUT
    widths(ws, [22, 34, 12, 8, 12, 8, 14, 12, 12, 12, 16])
    instructions(wb, "GST Invoice Template", [
        "1. Fill your business details at the top of the 'Invoice' sheet.",
        "2. Enter customer details. Put 2-digit state codes (e.g. 29 Karnataka, 27 Maharashtra, 07 Delhi, 33 Tamil Nadu).",
        "3. Enter items, quantity, rate and pick GST %. CGST/SGST/IGST split is automatic.",
        "4. Print to PDF and send. Duplicate the sheet for each new invoice and change the invoice number.",
        "Note: this is a calculation aid, not tax advice. Confirm GST rates with your CA."])
    wb.save(f"{OUT}/1_GST_Invoice.xlsx")

# 2 ---------------- Expense & profit tracker
def expenses():
    wb = Workbook(); ws = wb.active; ws.title = "Transactions"
    hdr(ws, 1, ["Date", "Type", "Category", "Description", "Amount (₹)", "Payment mode"])
    t_dv = DataValidation(type="list", formula1='"Income,Expense"'); ws.add_data_validation(t_dv)
    cats = ["Sales", "Rent", "Salaries", "Raw material", "Marketing", "Utilities", "Travel", "Software", "Taxes", "Other"]
    c_dv = DataValidation(type="list", formula1='"' + ",".join(cats) + '"'); ws.add_data_validation(c_dv)
    m_dv = DataValidation(type="list", formula1='"UPI,Cash,Card,Bank transfer"'); ws.add_data_validation(m_dv)
    for r in range(2, 502):
        for c in range(1, 7): ws.cell(row=r, column=c).fill = INPUT
        t_dv.add(ws.cell(row=r, column=2)); c_dv.add(ws.cell(row=r, column=3)); m_dv.add(ws.cell(row=r, column=6))
        ws.cell(row=r, column=1).number_format = "DD-MMM-YYYY"
    widths(ws, [14, 12, 16, 36, 14, 16])
    s = wb.create_sheet("Summary")
    hdr(s, 1, ["Category", "Total (₹)"])
    for i, c in enumerate(cats, 2):
        s.cell(row=i, column=1, value=c)
        s.cell(row=i, column=2, value=f'=SUMIFS(Transactions!$E:$E,Transactions!$C:$C,A{i})')
    n = len(cats) + 3
    for lbl, f, row in [("Total income", '=SUMIFS(Transactions!E:E,Transactions!B:B,"Income")', n),
                        ("Total expenses", '=SUMIFS(Transactions!E:E,Transactions!B:B,"Expense")', n + 1),
                        ("NET PROFIT", f"=B{n}-B{n+1}", n + 2),
                        ("Profit margin %", f'=IF(B{n}=0,0,B{n+2}/B{n})', n + 3)]:
        s.cell(row=row, column=1, value=lbl).font = Font(bold=True)
        s.cell(row=row, column=2, value=f).font = Font(bold=True)
    s.cell(row=n + 3, column=2).number_format = "0.0%"
    ch = BarChart(); ch.title = "Amount by category"; ch.height = 8; ch.width = 16
    ch.add_data(Reference(s, min_col=2, min_row=1, max_row=len(cats) + 1), titles_from_data=True)
    ch.set_categories(Reference(s, min_col=1, min_row=2, max_row=len(cats) + 1))
    s.add_chart(ch, "D2"); widths(s, [22, 16])
    instructions(wb, "Expense & Profit Tracker", [
        "1. Go to 'Transactions'. Add each sale (Income) and cost (Expense) in a new row.",
        "2. Pick Type, Category and Payment mode from the dropdowns.",
        "3. 'Summary' shows total income, expenses, net profit, margin and a chart. Nothing to calculate yourself.",
        "Tip: record every UPI payment the same day. 2 minutes a day saves hours at tax time."])
    wb.save(f"{OUT}/2_Expense_Profit_Tracker.xlsx")

# 3 ---------------- Stock tracker
def stock():
    wb = Workbook(); ws = wb.active; ws.title = "Stock"
    hdr(ws, 1, ["SKU", "Item", "Opening stock", "Purchased", "Sold", "Current stock", "Reorder level", "Status", "Cost price (₹)", "Selling price (₹)", "Stock value (₹)", "Profit/unit (₹)"])
    for r in range(2, 302):
        for c in (1, 2, 3, 4, 5, 7, 9, 10): ws.cell(row=r, column=c).fill = INPUT
        ws.cell(row=r, column=6, value=f'=IF(B{r}="","",C{r}+D{r}-E{r})')
        ws.cell(row=r, column=8, value=f'=IF(B{r}="","",IF(F{r}<=0,"OUT OF STOCK",IF(F{r}<=G{r},"REORDER NOW","OK")))')
        ws.cell(row=r, column=11, value=f'=IF(B{r}="","",F{r}*I{r})')
        ws.cell(row=r, column=12, value=f'=IF(B{r}="","",J{r}-I{r})')
    from openpyxl.formatting.rule import CellIsRule
    rng = "H2:H301"
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"OK"'], fill=PatternFill("solid", bgColor="C6EFCE")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"REORDER NOW"'], fill=PatternFill("solid", bgColor="FFEB9C")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"OUT OF STOCK"'], fill=PatternFill("solid", bgColor="FFC7CE")))
    ws.freeze_panes = "C2"; widths(ws, [10, 28, 12, 11, 9, 13, 13, 16, 13, 15, 15, 14])
    ws["N1"] = "Total stock value (₹)"; ws["N1"].font = Font(bold=True); ws["N2"] = "=SUM(K2:K301)"
    ws["N4"] = "Items to reorder"; ws["N4"].font = Font(bold=True); ws["N5"] = '=COUNTIF(H2:H301,"REORDER NOW")+COUNTIF(H2:H301,"OUT OF STOCK")'
    ws.column_dimensions["N"].width = 22
    instructions(wb, "Stock / Inventory Tracker", [
        "1. List each product: SKU, name, opening stock, cost price, selling price and reorder level.",
        "2. Update 'Purchased' and 'Sold' as stock moves. Current stock, status and value update automatically.",
        "3. Green = OK, Yellow = REORDER NOW, Red = OUT OF STOCK. Total stock value is in cell N2."])
    wb.save(f"{OUT}/3_Stock_Inventory_Tracker.xlsx")

# 4 ---------------- Client / lead tracker
def clients():
    wb = Workbook(); ws = wb.active; ws.title = "Leads"
    hdr(ws, 1, ["Name", "Phone", "Source", "Stage", "Quote (₹)", "Last contact", "Follow-up date", "Follow-up due?", "Notes"])
    st = DataValidation(type="list", formula1='"New,Contacted,Quoted,Negotiating,Won,Lost"'); ws.add_data_validation(st)
    for r in range(2, 402):
        for c in (1, 2, 3, 4, 5, 6, 7, 9): ws.cell(row=r, column=c).fill = INPUT
        st.add(ws.cell(row=r, column=4))
        ws.cell(row=r, column=8, value=f'=IF(OR(G{r}="",D{r}="Won",D{r}="Lost"),"",IF(G{r}<=TODAY(),"FOLLOW UP TODAY",""))')
        ws.cell(row=r, column=6).number_format = ws.cell(row=r, column=7).number_format = "DD-MMM-YYYY"
    widths(ws, [22, 14, 14, 14, 12, 14, 15, 18, 40])
    d = wb.create_sheet("Dashboard")
    hdr(d, 1, ["Stage", "Count", "Value (₹)"])
    for i, s in enumerate(["New", "Contacted", "Quoted", "Negotiating", "Won", "Lost"], 2):
        d.cell(row=i, column=1, value=s)
        d.cell(row=i, column=2, value=f'=COUNTIF(Leads!D:D,A{i})')
        d.cell(row=i, column=3, value=f'=SUMIFS(Leads!E:E,Leads!D:D,A{i})')
    d["A9"] = "Win rate"; d["B9"] = '=IF(B6+B7=0,0,B6/(B6+B7))'; d["B9"].number_format = "0%"
    d["A9"].font = Font(bold=True); widths(d, [18, 10, 16])
    instructions(wb, "Client & Lead Tracker", [
        "1. Add every enquiry on 'Leads'. Set stage from the dropdown and a follow-up date.",
        "2. The 'Follow-up due?' column flags people you must call today.",
        "3. 'Dashboard' shows how many leads per stage, rupee value, and your win rate.",
        "Most sales are lost by not following up. This sheet fixes that."])
    wb.save(f"{OUT}/4_Client_Lead_Tracker.xlsx")

if __name__ == "__main__":
    gst_invoice(); expenses(); stock(); clients()
    z = os.path.join(OUT, "Small_Business_Toolkit_India.zip")
    with zipfile.ZipFile(z, "w", zipfile.ZIP_DEFLATED) as zf:
        for f in sorted(os.listdir(OUT)):
            if f.endswith(".xlsx"): zf.write(os.path.join(OUT, f), f)
    print("built", z)
